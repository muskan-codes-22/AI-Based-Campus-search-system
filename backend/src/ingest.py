import hashlib
import logging
from pathlib import Path
from typing import List, Tuple

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

try:
    from .config import DATA_DIR, CHROMA_DIR, CHUNK_SIZE, CHUNK_OVERLAP
    from .vector_store import get_vector_store
except ImportError:
    from config import DATA_DIR, CHROMA_DIR, CHUNK_SIZE, CHUNK_OVERLAP
    from vector_store import get_vector_store

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".ppt",
    ".xlsx",
    ".xls",
    ".csv",
    ".txt",
    ".md",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stable_source_id(path: Path) -> str:
    try:
        relative = str(path.resolve().relative_to(DATA_DIR.resolve())).replace("\\", "/")
    except ValueError:
        relative = path.name
    return hashlib.sha256(relative.encode("utf-8")).hexdigest()[:24]


# ---------------------------------------------------------------------------
# Specialized Format Parsers
# ---------------------------------------------------------------------------

def _load_pdf(path: Path) -> List[Document]:
    """Extract pages from a PDF file."""
    pages = PyPDFLoader(str(path)).load()
    for idx, page in enumerate(pages):
        page.metadata["page_number"] = int(page.metadata.get("page", idx)) + 1
        page.metadata["source_filename"] = path.name
    return pages


def _load_docx(path: Path) -> List[Document]:
    """Extract structured paragraphs and Markdown tables from a Word document (.docx)."""
    import docx

    doc = docx.Document(str(path))
    sections: List[Document] = []
    current_text_blocks = []
    section_num = 1

    # Extract paragraphs and inline headings
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        # If heading style, start a new logical section/page
        if p.style.name.startswith("Heading") and current_text_blocks:
            full_text = "\n\n".join(current_text_blocks)
            sections.append(
                Document(
                    page_content=full_text,
                    metadata={"page_number": section_num, "source_filename": path.name},
                )
            )
            section_num += 1
            current_text_blocks = [f"## {txt}"]
        elif p.style.name.startswith("Heading"):
            current_text_blocks.append(f"## {txt}")
        else:
            current_text_blocks.append(txt)

    # Extract tables as formatted Markdown
    for table_idx, table in enumerate(doc.tables, start=1):
        rows_data = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows_data.append(row_cells)

        if rows_data:
            header = rows_data[0]
            separator = ["---"] * len(header)
            md_table_lines = [
                f"| {' | '.join(header)} |",
                f"| {' | '.join(separator)} |",
            ]
            for row in rows_data[1:]:
                md_table_lines.append(f"| {' | '.join(row)} |")

            table_md = "\n".join(md_table_lines)
            current_text_blocks.append(f"\n[Table {table_idx}]\n{table_md}\n")

    if current_text_blocks:
        full_text = "\n\n".join(current_text_blocks)
        sections.append(
            Document(
                page_content=full_text,
                metadata={"page_number": section_num, "source_filename": path.name},
            )
        )

    return sections


def _load_pptx(path: Path) -> List[Document]:
    """Extract slide titles, bullet points, and presenter notes from a PowerPoint (.pptx)."""
    import pptx

    prs = pptx.Presentation(str(path))
    slides_docs = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_lines = []
        slide_title = None

        # Extract shapes & text frames
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        if not slide_title and shape == slide.shapes.title:
                            slide_title = line
                            slide_lines.append(f"### Slide {slide_idx}: {slide_title}")
                        else:
                            slide_lines.append(f"- {line}")

        # Extract speaker notes if available
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_lines.append(f"\n[Presenter Notes: {notes}]")

        content = "\n".join(slide_lines).strip()
        if content:
            slides_docs.append(
                Document(
                    page_content=content,
                    metadata={"page_number": slide_idx, "source_filename": path.name},
                )
            )

    return slides_docs


def _load_excel(path: Path) -> List[Document]:
    """Extract sheets from Excel files (.xlsx, .xls) as Markdown tables."""
    import pandas as pd

    excel_file = pd.ExcelFile(str(path))
    sheet_docs = []

    for sheet_idx, sheet_name in enumerate(excel_file.sheet_names, start=1):
        df = pd.read_excel(excel_file, sheet_name=sheet_name)
        # Drop completely empty rows and columns
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if df.empty:
            continue

        try:
            table_md = df.to_markdown(index=False)
        except Exception:
            table_md = df.to_string(index=False)

        content = f"### Sheet: {sheet_name} (from {path.name})\n\n{table_md}"
        sheet_docs.append(
            Document(
                page_content=content,
                metadata={
                    "page_number": sheet_idx,
                    "sheet_name": sheet_name,
                    "source_filename": path.name,
                },
            )
        )

    return sheet_docs


def _load_csv(path: Path) -> List[Document]:
    """Extract CSV file as Markdown table."""
    import pandas as pd

    df = pd.read_csv(str(path))
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return []

    try:
        table_md = df.to_markdown(index=False)
    except Exception:
        table_md = df.to_string(index=False)

    content = f"### Dataset: {path.name}\n\n{table_md}"
    return [
        Document(
            page_content=content,
            metadata={"page_number": 1, "source_filename": path.name},
        )
    ]


def _load_text(path: Path) -> List[Document]:
    """Read plain text and markdown documents."""
    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            text = path.read_text(encoding=enc).strip()
            if text:
                return [
                    Document(
                        page_content=text,
                        metadata={"page_number": 1, "source_filename": path.name},
                    )
                ]
            return []
        except UnicodeDecodeError:
            continue
    return []


def load_documents_from_file(path: Path) -> List[Document]:
    """Router to load any supported document format."""
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _load_pdf(path)
    elif ext in {".docx", ".doc"}:
        return _load_docx(path)
    elif ext in {".pptx", ".ppt"}:
        return _load_pptx(path)
    elif ext in {".xlsx", ".xls"}:
        return _load_excel(path)
    elif ext == ".csv":
        return _load_csv(path)
    elif ext in {".txt", ".md"}:
        return _load_text(path)
    else:
        logger.warning(f"Unsupported file format: {path.suffix}")
        return []


# ---------------------------------------------------------------------------
# Universal Ingest Engine
# ---------------------------------------------------------------------------

def ingest_file(path: Path, store) -> Tuple[int, str]:
    """
    Ingest any supported document (PDF, Word, PPTX, Excel, CSV, Text) into Chroma.
    Tracks file hashes for automatic incremental change detection.
    """
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return 0, f"unsupported extension: {path.suffix}"

    file_hash = sha256_file(path)
    source_id = stable_source_id(path)

    existing = store.get(where={"source_id": source_id}, include=["metadatas"])
    old_metadata = existing.get("metadatas") or []

    if old_metadata and all(m.get("file_hash") == file_hash for m in old_metadata):
        logger.info(f"[skip] Unchanged: {path.name}")
        return len(old_metadata), "unchanged"

    # Remove old chunks if this document was previously indexed
    if old_metadata:
        store.delete(where={"source_id": source_id})
        logger.info(f"[update] Cleared previous index for: {path.name}")

    pages = load_documents_from_file(path)
    if not pages:
        logger.warning(f"[skip] No content extracted from: {path.name}")
        return 0, "empty"

    for idx, doc in enumerate(pages):
        doc.metadata["source_filename"] = path.name
        doc.metadata["source_id"] = source_id
        doc.metadata["file_hash"] = file_hash
        doc.metadata["file_type"] = path.suffix.lower().lstrip(".")
        if "page_number" not in doc.metadata:
            doc.metadata["page_number"] = idx + 1

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        add_start_index=True,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(pages)
    chunks = [c for c in chunks if c.page_content.strip()]

    if not chunks:
        logger.warning(f"[skip] No non-empty chunks in {path.name}")
        return 0, "empty"

    ids = []
    for index, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = index
        chunk.metadata["source_filename"] = path.name
        chunk.metadata["source_id"] = source_id
        chunk.metadata["file_hash"] = file_hash
        chunk.metadata["file_type"] = path.suffix.lower().lstrip(".")
        chunk.metadata["page_number"] = int(chunk.metadata.get("page_number", 1))

        chunk_id = hashlib.sha256(
            f"{source_id}:{file_hash}:{index}".encode("utf-8")
        ).hexdigest()
        chunk.metadata["chunk_id"] = chunk_id
        ids.append(chunk_id)

    store.add_documents(chunks, ids=ids)
    logger.info(f"[indexed] {path.name}: {len(pages)} sections/pages, {len(chunks)} chunks")
    return len(chunks), "indexed"


# Backwards compatibility alias
ingest_pdf = ingest_file


def delete_document_from_store(filename: str, store) -> int:
    """Delete chunks for a specific file by source_filename."""
    existing = store.get(where={"source_filename": filename}, include=["metadatas"])
    ids = existing.get("ids") or []
    if ids:
        store.delete(ids=ids)
        logger.info(f"[delete] Removed {len(ids)} chunks for: {filename}")
    return len(ids)


def main():
    files = [
        p for p in DATA_DIR.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    if not files:
        print(f"No supported documents found. Place files (.pdf, .docx, .pptx, .xlsx, .csv, .txt) in: {DATA_DIR}")
        return

    store = get_vector_store()
    total_chunks = 0
    for path in sorted(files):
        try:
            count, status = ingest_file(path, store)
            if status == "indexed":
                total_chunks += count
                print(f"[indexed] {path.name}: {count} chunks")
            elif status == "unchanged":
                print(f"[unchanged] {path.name}")
        except Exception as exc:
            print(f"[error] Could not ingest {path.name}: {exc}")

    print(f"\nDone. Total newly indexed chunks: {total_chunks}")
    print(f"Chroma storage directory: {CHROMA_DIR.resolve()}")


if __name__ == "__main__":
    main()
